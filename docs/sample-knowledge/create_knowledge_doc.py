"""One-off script to generate sample knowledge-base DOCX for AI sequences."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

OUT = Path(__file__).with_name("Coherent_Outreach_Company_Knowledge.docx")


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def build() -> None:
    doc = Document()

    title = doc.add_heading("Coherent Outreach — Company Knowledge Base", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT

    sub = doc.add_paragraph()
    run = sub.add_run("Internal document for sales, marketing, and AI-assisted sequence generation.")
    run.italic = True
    run.font.size = Pt(11)

    doc.add_paragraph()

    add_heading(doc, "1. Company Overview")
    doc.add_paragraph(
        "Coherent Outreach is a B2B growth company built on multi-channel outreach. "
        "We help revenue teams reach prospects where they actually respond — email, LinkedIn, "
        "phone, SMS, and WhatsApp — from one coordinated playbook instead of disconnected tools."
    )
    doc.add_paragraph(
        "Founded to solve the problem of fragmented follow-up, we combine sequencing, "
        "personalization, and reply intelligence so every touch feels intentional, not spammy."
    )

    add_heading(doc, "2. Mission & Value Proposition", level=2)
    add_bullets(
        doc,
        [
            "Mission: Make outbound feel human at scale by orchestrating the right message on the right channel at the right time.",
            "Value proposition: One platform to design, run, and optimize multi-step outreach across email and social — with AI that learns from replies.",
            "Tagline options: “Outreach that stays coherent.” / “Every channel. One story.” / “Stop guessing — start sequencing.”",
        ],
    )

    add_heading(doc, "3. What We Sell", level=2)
    add_bullets(
        doc,
        [
            "Multi-channel sequence builder — drag-and-drop steps across email, LinkedIn connect/DM, call tasks, SMS, and WhatsApp.",
            "AI sequence generation — turn a brief and uploaded knowledge into ready-to-edit campaigns in minutes.",
            "Reply intelligence — sentiment labeling, hot-lead surfacing, and AI-drafted follow-ups grounded in conversation history.",
            "Lead vault (RAG) — per-lead context so follow-ups reference what was actually said, not generic templates.",
            "Deliverability-aware sending — daily caps, jitter, and channel-specific limits to protect sender reputation.",
        ],
    )

    add_heading(doc, "4. Ideal Customer Profile (ICP)", level=2)
    add_bullets(
        doc,
        [
            "B2B companies with 5–200 person sales or growth teams.",
            "Founders, SDR leaders, and RevOps at SaaS, agencies, and professional services firms.",
            "Teams already doing outbound but struggling with low reply rates, tool sprawl, or inconsistent messaging.",
            "Geography: North America, UK, and India-first SaaS; open to global English-speaking markets.",
            "Deal size sweet spot: $3k–$50k ACV; sales cycle 2–8 weeks when champion is Head of Sales or Growth.",
        ],
    )

    add_heading(doc, "5. Pain Points We Solve", level=2)
    add_bullets(
        doc,
        [
            "Prospects ignore single-channel email blasts — we coordinate email + LinkedIn + timely call/SMS nudges.",
            "Reps rewrite the same follow-ups manually — AI drafts use vault context and approved company messaging.",
            "No visibility into which sequence steps actually convert — dashboard shows sentiment, hot leads, and at-risk enrolments.",
            "LinkedIn and email live in separate tools — one sequence timeline across channels.",
            "Compliance and fatigue — built-in caps, unsubscribe handling, and suppression across all channels.",
        ],
    )

    add_heading(doc, "6. Multi-Channel Playbook", level=2)

    add_heading(doc, "Email", level=3)
    add_bullets(
        doc,
        [
            "Best for: first touch, value-led follow-ups, sharing one-pagers or case studies.",
            "Keep subject lines under 7 words; lead with their outcome, not our product name.",
            "Recommended cadence: Day 0 intro, Day 3 bump, Day 7 new angle, Day 12 break-up or soft close.",
        ],
    )

    add_heading(doc, "LinkedIn", level=3)
    add_bullets(
        doc,
        [
            "Best for: warm-up before email, mutual connections, and short DM follow-ups after email opens.",
            "Connect note: mention a specific trigger (funding, hiring, product launch) — never pitch in the connect request.",
            "DM follow-up: reference the email subject line; ask one clear question.",
        ],
    )

    add_heading(doc, "Phone / Call tasks", level=3)
    add_bullets(
        doc,
        [
            "Best for: high-intent leads (positive reply, pricing question, demo request).",
            "Call goal: confirm fit in 15 minutes, not a full demo on first call.",
            "Voicemail script: name, company, one reason you called, and a specific callback window.",
        ],
    )

    add_heading(doc, "SMS & WhatsApp", level=3)
    add_bullets(
        doc,
        [
            "Best for: time-sensitive reminders, meeting confirmations, and re-engaging stalled threads.",
            "Always opt-in friendly; keep messages under 320 characters; include an easy opt-out.",
            "WhatsApp: use for regions where business WhatsApp is normal (India, UK, EU, LATAM).",
        ],
    )

    add_heading(doc, "7. Messaging Pillars", level=2)
    pillars = [
        ("Coherence", "Every touch should reference the same core story — problem, proof, next step."),
        ("Proof", "Lead with customer outcomes: reply rate lift, time saved, fewer tools."),
        ("Speed to value", "Teams launch a 5-step sequence in under an hour with AI + knowledge upload."),
        ("Safety", "Reputation matters — caps, jitter, and sentiment-aware pauses protect domains and LinkedIn accounts."),
    ]
    for name, desc in pillars:
        p = doc.add_paragraph()
        p.add_run(f"{name}: ").bold = True
        p.add_run(desc)

    add_heading(doc, "8. Tone & Voice Guidelines", level=2)
    add_bullets(
        doc,
        [
            "Professional but conversational — write like a helpful peer, not a billboard.",
            "Confident, not arrogant; specific, not buzzword-heavy.",
            "Avoid: “synergy,” “revolutionary,” “game-changing,” “hope this finds you well.”",
            "Prefer: concrete numbers, named outcomes, and one clear CTA per message.",
            "Personalization tokens: {{first_name}}, {{company}}, {{title}}, {{trigger_event}}.",
        ],
    )

    add_heading(doc, "9. Approved CTAs", level=2)
    add_bullets(
        doc,
        [
            "Worth a 15-minute call this week to see if multi-channel sequencing fits your stack?",
            "Open to a quick walkthrough of how teams run email + LinkedIn from one timeline?",
            "Should I send a 2-minute loom on how we draft follow-ups from actual reply context?",
            "If timing is bad, happy to circle back next quarter — which month works better?",
            "Reply “demo” and I’ll send three calendar slots.",
        ],
    )

    add_heading(doc, "10. Objection Handling", level=2)
    objections = [
        (
            "We already use [Salesloft / Outreach / HubSpot sequences]",
            "Those tools excel at email. We unify email, LinkedIn, SMS, and call tasks in one sequence with reply-aware AI follow-ups — many teams keep HubSpot as CRM and use us for orchestration.",
        ),
        (
            "We don’t want to spam LinkedIn",
            "Daily caps and human-in-the-loop steps keep volume safe. Connect notes are trigger-based, not bulk pitch blasts.",
        ),
        (
            "AI content feels generic",
            "Upload your positioning doc (like this one) — generation and follow-up drafts pull from your knowledge base and per-lead vault, not blank prompts.",
        ),
        (
            "No budget this quarter",
            "Understood. Can I send a one-page ROI snapshot based on your team size? No call required.",
        ),
    ]
    for obj, response in objections:
        p = doc.add_paragraph()
        p.add_run(f"Objection: {obj}").bold = True
        doc.add_paragraph(f"Response: {response}")

    add_heading(doc, "11. Proof Points & Social Proof", level=2)
    add_bullets(
        doc,
        [
            "Customers report 2–3× reply rates when adding LinkedIn steps to existing email sequences.",
            "Average time to first live sequence: under 60 minutes including knowledge upload.",
            "Sentiment tagging surfaces “hot” replies (pricing, timeline, demo interest) on the dashboard within minutes.",
            "Example win: B2B SaaS team (40 SDRs) consolidated 4 tools into one workflow and cut follow-up drafting time by ~70%.",
        ],
    )

    add_heading(doc, "12. Pricing Positioning (High Level)", level=2)
    doc.add_paragraph(
        "We price per seat with volume tiers for sequences and AI usage. "
        "Position as mid-market: more capable than single-channel sequencers, "
        "more affordable than stitching together five point tools. "
        "Offer a 14-day pilot for teams over 5 seats. "
        "Do not quote exact dollar amounts in cold outreach — invite a scoping call."
    )

    add_heading(doc, "13. Compliance & Opt-Out", level=2)
    add_bullets(
        doc,
        [
            "Honor unsubscribe and “not interested” across all channels immediately.",
            "CAN-SPAM: physical address and clear unsubscribe in marketing emails where required.",
            "LinkedIn: respect platform limits; no automation that violates LinkedIn Terms of Service.",
            "SMS/WhatsApp: obtain appropriate consent; include STOP language where applicable.",
        ],
    )

    add_heading(doc, "14. Sample Sequence Angles (for AI generation)", level=2)
    angles = [
        "SaaS founder — book demo: 3 email + 1 LinkedIn connect + 1 call task over 10 days.",
        "SDR leader — tool consolidation: lead with pain of 4 tools, offer loom walkthrough.",
        "Agency owner — client results: case study on reply rate lift with multi-channel.",
        "RevOps — integration story: works alongside CRM; focus on timeline + sentiment dashboard.",
    ]
    for angle in angles:
        doc.add_paragraph(angle, style="List Number")

    doc.add_paragraph()
    footer = doc.add_paragraph()
    footer.add_run("Document version: 1.0 | Last updated: June 2026 | Owner: Marketing & Sales Enablement").italic = True

    doc.save(OUT)
    print(f"Created: {OUT}")


if __name__ == "__main__":
    build()
