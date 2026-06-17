"""Parse, chunk, embed, and index documents for sequence RAG."""
from __future__ import annotations

import asyncio
import logging
from dataclasses import dataclass
from pathlib import Path
from uuid import uuid4

from fastapi import HTTPException, UploadFile
from langchain_text_splitters import RecursiveCharacterTextSplitter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from outreach.config import get_settings
from outreach.models.sequence_rag_document import SequenceRagDocument
from outreach.schemas.sequence_rag import SequenceRagDocumentOut
from outreach.services import embedding_service, qdrant_store

log = logging.getLogger("outreach.document_ingest")

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}
MAX_DOC_BYTES = 20 * 1024 * 1024
MIN_EXTRACT_CHARS = 50


class DocParseError(Exception):
    pass


@dataclass(slots=True)
class ChunkMeta:
    text: str
    chunk_index: int
    total_chunks: int
    char_start: int
    char_end: int


def _uploads_dir(user_id: int) -> Path:
    base = get_settings().vault_dir.parent / "uploads" / str(user_id) / "sequence_rag"
    base.mkdir(parents=True, exist_ok=True)
    return base


def _file_type(ext: str) -> str:
    return ext.lstrip(".").lower()


def parse_document(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        parts = []
        for page in reader.pages:
            parts.append(page.extract_text() or "")
        return "\n".join(parts)
    if ext == ".docx":
        from docx import Document

        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    if ext in (".txt", ".md"):
        return path.read_text(encoding="utf-8", errors="replace")
    raise DocParseError(f"unsupported extension: {ext}")


def recursive_chunk(text: str, chunk_size: int | None = None, chunk_overlap: int | None = None) -> list[ChunkMeta]:
    settings = get_settings()
    size = chunk_size or settings.rag_chunk_size
    overlap = chunk_overlap or settings.rag_chunk_overlap
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=size,
        chunk_overlap=overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    pieces = splitter.split_text(text)
    total = len(pieces)
    metas: list[ChunkMeta] = []
    search_from = 0
    for i, piece in enumerate(pieces):
        start = text.find(piece, search_from)
        if start == -1:
            start = search_from
        end = start + len(piece)
        search_from = max(0, end - overlap)
        metas.append(
            ChunkMeta(
                text=piece,
                chunk_index=i,
                total_chunks=total,
                char_start=start,
                char_end=end,
            )
        )
    return metas


def _to_out(doc: SequenceRagDocument) -> SequenceRagDocumentOut:
    return SequenceRagDocumentOut(
        id=doc.id,
        filename=doc.filename,
        mime_type=doc.mime_type,
        file_size=doc.file_size,
        status=doc.status,  # type: ignore[arg-type]
        chunk_count=doc.chunk_count,
        error_message=doc.error_message,
        created_at=doc.created_at,
        updated_at=doc.updated_at,
    )


async def ingest_document(
    session: AsyncSession,
    user_id: int,
    upload: UploadFile,
) -> SequenceRagDocumentOut:
    if not upload.filename:
        raise HTTPException(400, "missing filename")
    ext = Path(upload.filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(400, f"supported formats: {', '.join(sorted(ALLOWED_EXTENSIONS))}")

    content = await upload.read()
    if len(content) > MAX_DOC_BYTES:
        raise HTTPException(413, f"file too large (max {MAX_DOC_BYTES // (1024 * 1024)} MB)")
    if not content:
        raise HTTPException(400, "empty file")

    token = uuid4().hex
    out_path = _uploads_dir(user_id) / f"{token}{ext}"
    out_path.write_bytes(content)

    doc = SequenceRagDocument(
        user_id=user_id,
        filename=upload.filename,
        storage_path=str(out_path),
        mime_type=upload.content_type or "",
        file_size=len(content),
        status="pending",
    )
    session.add(doc)
    await session.commit()
    await session.refresh(doc)

    if not qdrant_store.is_available():
        doc.status = "failed"
        doc.error_message = "Qdrant unavailable — start with: docker compose up qdrant"
        await session.commit()
        await session.refresh(doc)
        raise HTTPException(503, doc.error_message)

    try:
        text = await asyncio.to_thread(parse_document, out_path)
        if len(text.strip()) < MIN_EXTRACT_CHARS:
            raise DocParseError(
                "no extractable text (scanned PDF or empty document?)"
            )
        chunks = recursive_chunk(text)
        if not chunks:
            raise DocParseError("document produced zero chunks")

        qdrant_store.delete_document_chunks(user_id, doc.id)

        payloads = [
            qdrant_store.ChunkPayload(
                text=c.text,
                chunk_index=c.chunk_index,
                total_chunks=c.total_chunks,
                char_start=c.char_start,
                char_end=c.char_end,
                filename=doc.filename,
                file_type=_file_type(ext),
            )
            for c in chunks
        ]
        vectors = await asyncio.to_thread(
            embedding_service.embed_texts,
            [p.text for p in payloads],
        )
        await asyncio.to_thread(
            qdrant_store.upsert_chunks,
            user_id,
            doc.id,
            payloads,
            vectors,
        )

        doc.status = "indexed"
        doc.chunk_count = len(chunks)
        doc.error_message = None
    except DocParseError as e:
        doc.status = "failed"
        doc.error_message = str(e)
        log.warning("document ingest failed doc=%s: %s", doc.id, e)
    except Exception as e:  # noqa: BLE001
        doc.status = "failed"
        doc.error_message = f"{type(e).__name__}: {e}"
        log.exception("document ingest error doc=%s", doc.id)

    await session.commit()
    await session.refresh(doc)
    return _to_out(doc)


async def list_documents(session: AsyncSession, user_id: int) -> list[SequenceRagDocumentOut]:
    rows = (await session.execute(
        select(SequenceRagDocument)
        .where(SequenceRagDocument.user_id == user_id)
        .order_by(SequenceRagDocument.id.desc())
    )).scalars().all()
    return [_to_out(r) for r in rows]


async def delete_document(session: AsyncSession, user_id: int, doc_id: int) -> bool:
    doc = await session.scalar(
        select(SequenceRagDocument).where(
            SequenceRagDocument.id == doc_id,
            SequenceRagDocument.user_id == user_id,
        )
    )
    if doc is None:
        return False
    try:
        await asyncio.to_thread(qdrant_store.delete_document_chunks, user_id, doc_id)
    except Exception as e:  # noqa: BLE001
        log.warning("qdrant delete failed for doc %s: %s", doc_id, e)
    try:
        Path(doc.storage_path).unlink(missing_ok=True)
    except OSError as e:
        log.warning("file delete failed for doc %s: %s", doc_id, e)
    await session.delete(doc)
    await session.commit()
    return True


async def resolve_document_ids(
    session: AsyncSession,
    user_id: int,
    document_ids: list[int] | None,
) -> list[int]:
    """Return indexed document ids for this user. Empty list input → all indexed docs."""
    q = select(SequenceRagDocument.id).where(
        SequenceRagDocument.user_id == user_id,
        SequenceRagDocument.status == "indexed",
    )
    if document_ids:
        q = q.where(SequenceRagDocument.id.in_(document_ids))
    return list((await session.execute(q)).scalars().all())
